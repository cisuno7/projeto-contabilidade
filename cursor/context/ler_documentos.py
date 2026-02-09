#!/usr/bin/env python3
"""Script para extrair conteúdo dos documentos da pasta context"""

import os
import sys

BASE = os.path.dirname(os.path.abspath(__file__))

def ler_xlsx(caminho, max_linhas=30):
    """Lê planilha Excel e retorna conteúdo"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        resultado = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            linhas = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_linhas:
                    break
                linhas.append([str(c) if c is not None else "" for c in row])
            resultado.append((sheet_name, linhas))
        wb.close()
        return resultado
    except Exception as e:
        return [("ERRO", str(e))]

def ler_docx(caminho):
    """Lê documento Word"""
    try:
        import docx
        doc = docx.Document(caminho)
        return [p.text for p in doc.paragraphs if p.text.strip()]
    except Exception as e:
        return [f"ERRO: {e}"]

def ler_pdf(caminho, max_paginas=5):
    """Lê PDF"""
    try:
        import pdfplumber
        texto = []
        with pdfplumber.open(caminho) as pdf:
            for i, page in enumerate(pdf.pages):
                if i >= max_paginas:
                    break
                text = page.extract_text()
                if text:
                    texto.append(f"--- Página {i+1} ---")
                    texto.append(text[:3000])  # primeiros 3000 chars por página
        return "\n".join(texto) if texto else "Nenhum texto extraído"
    except ImportError:
        try:
            import PyPDF2
            with open(caminho, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                texto = []
                for i in range(min(5, len(reader.pages))):
                    page = reader.pages[i]
                    texto.append(page.extract_text() or "")
            return "\n".join(texto)
        except Exception as e:
            return f"ERRO: {e}"
    except Exception as e:
        return f"ERRO: {e}"

def main():
    print("=" * 60)
    print("LEITURA DOS DOCUMENTOS - cursor/context")
    print("=" * 60)

    # 1. NCM.xlsx
    print("\n## 1. NCM.xlsx (Planilha Padrão Nacional)")
    print("-" * 40)
    ncm_path = os.path.join(BASE, "NCM.xlsx")
    if os.path.exists(ncm_path):
        for sheet_name, linhas in ler_xlsx(ncm_path):
            print(f"\nAba: {sheet_name}")
            for i, row in enumerate(linhas):
                print(f"  {i+1}: {row}")
    else:
        print("Arquivo não encontrado")

    # 2. Planilha Cadastro Padaria
    print("\n\n## 2. Planilha-Cadastro_Produto-Padaria-Brasil-MEJ Padaria modificado.xlsx")
    print("-" * 40)
    padaria_path = os.path.join(BASE, "Planilha-Cadastro_Produto-Padaria-Brasil-MEJ Padaria modificado.xlsx")
    if os.path.exists(padaria_path):
        for sheet_name, linhas in ler_xlsx(padaria_path):
            print(f"\nAba: {sheet_name}")
            for i, row in enumerate(linhas):
                print(f"  {i+1}: {row}")
    else:
        print("Arquivo não encontrado")

    # 3. CSOSN.docx
    print("\n\n## 3. CSOSN.docx")
    print("-" * 40)
    csosn_path = os.path.join(BASE, "CSOSN.docx")
    if os.path.exists(csosn_path):
        paragrafos = ler_docx(csosn_path)
        for p in paragrafos:
            print(p)
    else:
        print("Arquivo não encontrado")

    # 4. CEST SP.pdf
    print("\n\n## 4. CEST SP.pdf")
    print("-" * 40)
    cest_path = os.path.join(BASE, "CEST SP.pdf")
    if os.path.exists(cest_path):
        conteudo = ler_pdf(cest_path)
        print(conteudo[:5000] if len(conteudo) > 5000 else conteudo)
    else:
        print("Arquivo não encontrado")

    # 5. Codigos CST PIS COFINS
    print("\n\n## 5. Codigos_CST_PIS_COFINS_260205_105256.pdf")
    print("-" * 40)
    cst_path = os.path.join(BASE, "Codigos_CST_PIS_COFINS_260205_105256.pdf")
    if os.path.exists(cst_path):
        conteudo = ler_pdf(cst_path)
        print(conteudo[:5000] if len(conteudo) > 5000 else conteudo)
    else:
        print("Arquivo não encontrado")

    print("\n" + "=" * 60)
    print("FIM DA LEITURA")
    print("=" * 60)

if __name__ == "__main__":
    main()
